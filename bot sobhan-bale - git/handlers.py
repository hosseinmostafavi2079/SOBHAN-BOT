import os
import re
import asyncio
import pandas as pd
import jdatetime  # برای تاریخ شمسی
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from telegram import Update, ReplyKeyboardMarkup, InlineKeyboardMarkup, InlineKeyboardButton
from telegram.ext import ContextTypes, ConversationHandler
from telegram.error import BadRequest

import config
from config import *
from utils import Utils, Security, PDF_SUPPORT
from wc_layer import WCService

# --- سیستم قفل ---
async def check_is_busy(update, context):
    if context.user_data.get('is_busy', False):
        if update.callback_query: await update.callback_query.answer("⚠️ لطفاً صبر کنید...", show_alert=True)
        else: await update.message.reply_text("⏳ ربات مشغول است...")
        return True
    return False

def set_busy(context, state: bool):
    context.user_data['is_busy'] = state

async def safe_reply(update, text, markup=None, parse_mode=None):
    try:
        if update.message: return await update.message.reply_text(text, reply_markup=markup, parse_mode=parse_mode)
        elif update.callback_query: return await update.callback_query.message.reply_text(text, reply_markup=markup, parse_mode=parse_mode)
    except: pass

async def safe_edit(context, chat_id, msg_id, text, markup=None, parse_mode=None):
    try: await context.bot.edit_message_text(chat_id=chat_id, message_id=msg_id, text=text, reply_markup=markup, parse_mode=parse_mode)
    except BadRequest: pass
    except Exception as e: config.logger.error(f"Edit Error: {e}")

# --- کیبوردها ---
def get_main_menu(user_id):
    user = config.AUTHORIZED_USERS.get(user_id, {})
    role = user.get('role', 'manager')
    btns = [["🔍 استعلام کالا", "📋 لیست محصولات"], ["📈 تغییر قیمت گروهی", "🗑 حذف قیمت گروهی"]]
    if role == "admin": btns.append(["📥 خروجی گرفتن", "📝 گزارش‌ها"])
    else: btns.append(["📥 خروجی گرفتن"])
    return ReplyKeyboardMarkup(btns, resize_keyboard=True)

CANCEL_KB = ReplyKeyboardMarkup([["انصراف"]], resize_keyboard=True)

def get_product_kb(pid):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("💰 تغییر قیمت", callback_data=f"edit_price_{pid}"), InlineKeyboardButton("🗑 حذف قیمت", callback_data=f"remove_price_{pid}")],
        [InlineKeyboardButton("➕ افزایش", callback_data=f"add_stock_{pid}"), InlineKeyboardButton("➖ کاهش", callback_data=f"reduce_stock_{pid}")],
        [InlineKeyboardButton("✅ موجود", callback_data=f"set_instock_{pid}"), InlineKeyboardButton("❌ ناموجود", callback_data=f"set_outstock_{pid}")],
        [InlineKeyboardButton("🔄 بروزرسانی", callback_data=f"refresh_{pid}")]
    ])

# --- هندلرها ---
@Security.authorize
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await safe_reply(update, "👋 سلام. به پنل مدیریت خوش آمدید.", get_main_menu(update.effective_user.id))
    return MENU_STATE

async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    set_busy(context, False)
    await safe_reply(update, "❌ لغو شد.", get_main_menu(update.effective_user.id))
    return MENU_STATE

async def deep_link_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    match = re.search(r"/info_?(.+)", update.message.text)
    if match:
        if await check_is_busy(update, context): return
        await show_product(update, context, sku=match.group(1).split()[0])

async def menu_router(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    if await check_is_busy(update, context): return MENU_STATE
    
    user_id = update.effective_user.id
    user_role = config.AUTHORIZED_USERS.get(user_id, {}).get('role')

    if txt == "🔍 استعلام کالا": await safe_reply(update, "🔖 SKU را وارد کنید:", CANCEL_KB); return MENU_STATE
    elif txt == "📋 لیست محصولات": return await list_start(update, context)
    elif txt == "📈 تغییر قیمت گروهی": return await bulk_start(update, context)
    elif txt == "🗑 حذف قیمت گروهی": return await remove_start(update, context)
    elif txt == "📥 خروجی گرفتن": return await export_start(update, context)
    elif txt == "📝 گزارش‌ها":
        if user_role == "admin":
            if os.path.exists(config.LOG_FILE): await update.message.reply_document(open(config.LOG_FILE, 'rb'))
            else: await safe_reply(update, "لاگی موجود نیست.")
        return MENU_STATE
    elif txt == "انصراف": return await cancel(update, context)

    await show_product(update, context, sku=Utils.to_english(txt))
    return MENU_STATE

async def show_product(update, context, sku=None, pid=None, msg_id_to_edit=None):
    if sku:
        temp = await safe_reply(update, "🔎 جستجو...")
        product = await WCService.get_product(sku=sku)
        if temp: await context.bot.delete_message(update.effective_chat.id, temp.message_id)
    elif pid:
        product = await WCService.get_product(pid=pid)
    else: return

    if not product:
        if update.message: await safe_reply(update, "❌ یافت نشد.", get_main_menu(update.effective_user.id))
        return

    real_sku = product.get('sku') or str(product['id'])
    context.user_data['active_sku'] = real_sku

    p_id = product['id']
    name = product['name']
    stock = "✅" if product['stock_status'] == 'instock' else "🔴"
    qty = product.get('stock_quantity') or 0
    price = Utils.format_price(product.get('regular_price'))
    sale = Utils.format_price(product.get('sale_price'))
    
    txt = (f"📦 <b>{name}</b>\nSKU: <code>{real_sku}</code>\n"
           f"وضعیت: {stock} ({qty})\nقیمت: {price} | حراج: {sale}\n"
           f"<a href='{product['permalink']}'>لینک</a>")
    
    kb = get_product_kb(p_id)
    if msg_id_to_edit: await safe_edit(context, update.effective_chat.id, msg_id_to_edit, txt, kb, 'HTML')
    else: await safe_reply(update, txt, kb, 'HTML')

async def callback_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer("🔄")
    try: action, pid_str = query.data.rsplit('_', 1); pid = int(pid_str)
    except: return
    
    context.user_data['active_pid'] = pid
    sku_log = context.user_data.get('active_sku', str(pid))
    user_id = update.effective_user.id

    if action == "refresh": await show_product(update, context, pid=pid, msg_id_to_edit=query.message.message_id)
    elif action == "edit_price": await safe_reply(update, "💰 قیمت جدید:", CANCEL_KB); return WAIT_INPUT_PRICE
    elif action == "add_stock": await safe_reply(update, "➕ افزودن:", CANCEL_KB); return WAIT_INPUT_ADD
    elif action == "reduce_stock": await safe_reply(update, "➖ کاهش:", CANCEL_KB); return WAIT_INPUT_REDUCE
    
    elif action == "remove_price":
        await WCService.update_product(pid, {"regular_price": "", "sale_price": ""})
        Security.log_audit(user_id, "حذف قیمت", sku_log, "پاکسازی")
        await show_product(update, context, pid=pid, msg_id_to_edit=query.message.message_id)
        
    elif action == "set_instock":
        await WCService.update_product(pid, {"stock_status": "instock", "manage_stock": True})
        Security.log_audit(user_id, "موجود کردن", sku_log, "InStock")
        await show_product(update, context, pid=pid, msg_id_to_edit=query.message.message_id)
        
    elif action == "set_outstock":
        await WCService.update_product(pid, {"stock_status": "outofstock", "manage_stock": True, "stock_quantity": 0})
        Security.log_audit(user_id, "ناموجود کردن", sku_log, "OutStock")
        await show_product(update, context, pid=pid, msg_id_to_edit=query.message.message_id)
    return MENU_STATE

# --- Inputs ---
async def input_price(update, context):
    txt = Utils.to_english(update.message.text)
    if txt == "انصراف": return await cancel(update, context)
    if not txt.isdigit(): await safe_reply(update, "❌ عدد وارد کنید."); return WAIT_INPUT_PRICE
    
    pid = context.user_data['active_pid']
    sku_log = context.user_data.get('active_sku', str(pid))
    
    await WCService.update_product(pid, {"regular_price": txt})
    Security.log_audit(update.effective_user.id, "تغییر قیمت", sku_log, txt)
    
    await safe_reply(update, "✅ انجام شد.", get_main_menu(update.effective_user.id))
    await show_product(update, context, pid=pid)
    return MENU_STATE

async def input_add_stock(update, context):
    txt = Utils.to_english(update.message.text)
    if txt == "انصراف": return await cancel(update, context)
    if not txt.isdigit(): return WAIT_INPUT_ADD
    
    pid = context.user_data['active_pid']
    sku_log = context.user_data.get('active_sku', str(pid))
    p = await WCService.get_product(pid=pid)
    new_qty = (p.get('stock_quantity') or 0) + int(txt)
    
    await WCService.update_product(pid, {"stock_quantity": new_qty, "manage_stock": True, "stock_status": "instock"})
    Security.log_audit(update.effective_user.id, "افزایش موجودی", sku_log, f"+{txt}")
    
    await safe_reply(update, f"✅ موجودی جدید: {new_qty}", get_main_menu(update.effective_user.id))
    await show_product(update, context, pid=pid)
    return MENU_STATE

async def input_reduce_stock(update, context):
    txt = Utils.to_english(update.message.text)
    if txt == "انصراف": return await cancel(update, context)
    if not txt.isdigit(): return WAIT_INPUT_REDUCE
    
    pid = context.user_data['active_pid']
    sku_log = context.user_data.get('active_sku', str(pid))
    p = await WCService.get_product(pid=pid)
    new_qty = max(0, (p.get('stock_quantity') or 0) - int(txt))
    
    await WCService.update_product(pid, {"stock_quantity": new_qty, "manage_stock": True})
    Security.log_audit(update.effective_user.id, "کاهش موجودی", sku_log, f"-{txt}")
    
    await safe_reply(update, f"✅ موجودی جدید: {new_qty}", get_main_menu(update.effective_user.id))
    await show_product(update, context, pid=pid)
    return MENU_STATE

# --- List Module ---
async def list_start(update, context):
    await safe_reply(update, "فیلتر؟", ReplyKeyboardMarkup([["همه"], ["موجود", "ناموجود"], ["انصراف"]], resize_keyboard=True))
    return WAIT_LIST_FILTER

async def list_filter(update, context):
    txt = update.message.text
    if txt == "انصراف": return await cancel(update, context)
    context.user_data['list_filter'] = "instock" if "موجود" in txt and "نا" not in txt else ("outofstock" if "ناموجود" in txt else None)
    
    msg = await safe_reply(update, "⏳ دریافت دسته‌ها...")
    cats = await WCService.get_categories(0)
    await context.bot.delete_message(update.effective_chat.id, msg.message_id)
    
    btns = [[c['name']] for c in cats] + [["همه محصولات"], ["انصراف"]]
    context.user_data['cat_map'] = {c['name']: c['id'] for c in cats}
    await safe_reply(update, "📂 انتخاب دسته:", ReplyKeyboardMarkup(btns, resize_keyboard=True))
    return WAIT_LIST_PARENT

async def list_parent(update, context):
    txt = update.message.text
    if txt == "انصراف": return await cancel(update, context)
    if txt == "همه محصولات": return await execute_list(update, context, None)
    pid = context.user_data['cat_map'].get(txt)
    if not pid:
        await safe_reply(update, "⚠️ نامعتبر.")
        return WAIT_LIST_PARENT
    subs = await WCService.get_categories(pid)
    if not subs: return await execute_list(update, context, pid)
    context.user_data['sub_map'] = {c['name']: c['id'] for c in subs}
    context.user_data['parent_id'] = pid
    btns = [[c['name']] for c in subs] + [[f"همه {txt}"], ["انصراف"]]
    await safe_reply(update, "📂 زیردسته:", ReplyKeyboardMarkup(btns, resize_keyboard=True))
    return WAIT_LIST_SUB

async def list_sub(update, context):
    txt = update.message.text
    if txt == "انصراف": return await cancel(update, context)
    cat_id = context.user_data['sub_map'].get(txt) or context.user_data.get('parent_id')
    if not cat_id and txt.startswith("همه "): cat_id = context.user_data.get('parent_id')
    return await execute_list(update, context, cat_id)

async def execute_list(update, context, cat_id):
    if await check_is_busy(update, context): return MENU_STATE
    set_busy(context, True)
    msg = await safe_reply(update, "🚀 دریافت لیست...")
    buffer, count = "", 0
    status = context.user_data.get('list_filter')
    try:
        async for products in WCService.fetch_products_generator(cat_id, status):
            if count % 20 == 0: await safe_edit(context, update.effective_chat.id, msg.message_id, f"📥 دریافت {count} محصول...")
            for p in products:
                icon = "✅" if p['stock_status'] == 'instock' else "🔴"
                row = f"{icon} {p['name']}\n🆔 /info_{p['sku']} | 💰 {Utils.format_price(p.get('regular_price'))}\n➖➖➖\n"
                if len(buffer) + len(row) > 4000:
                    await safe_reply(update, buffer, parse_mode='HTML')
                    buffer = ""
                buffer += row
                count += 1
        try: await context.bot.delete_message(update.effective_chat.id, msg.message_id)
        except: pass
        if buffer: await safe_reply(update, buffer, parse_mode='HTML')
        await safe_reply(update, f"🏁 پایان ({count} محصول).", get_main_menu(update.effective_user.id))
    finally: set_busy(context, False)
    return MENU_STATE

# --- Bulk Module ---
async def bulk_start(update, context):
    await safe_reply(update, "دامنه:", ReplyKeyboardMarkup([["همه محصولات"], ["انتخاب دسته‌بندی"], ["انصراف"]], resize_keyboard=True))
    return WAIT_BULK_SCOPE

async def bulk_scope(update, context):
    txt = update.message.text
    if txt == "انصراف": return await cancel(update, context)
    if txt == "همه محصولات": context.user_data['bulk_cat'] = None; return await bulk_ask_method(update, context)
    msg = await safe_reply(update, "⏳...")
    cats = await WCService.get_categories(0)
    await context.bot.delete_message(update.effective_chat.id, msg.message_id)
    context.user_data['cat_map'] = {c['name']: c['id'] for c in cats}
    await safe_reply(update, "دسته:", ReplyKeyboardMarkup([[c['name']] for c in cats] + [["انصراف"]], resize_keyboard=True))
    return WAIT_BULK_CAT

async def bulk_cat(update, context):
    context.user_data['bulk_cat'] = context.user_data['cat_map'].get(update.message.text)
    if not context.user_data['bulk_cat']: return WAIT_BULK_CAT
    context.user_data['bulk_name'] = update.message.text
    return await bulk_ask_method(update, context)

async def bulk_ask_method(update, context):
    await safe_reply(update, "نوع:", ReplyKeyboardMarkup([["مبلغ ثابت (تومان)"], ["درصدی (%)"]], resize_keyboard=True))
    return WAIT_BULK_METHOD

async def bulk_method(update, context):
    context.user_data['bulk_type'] = 'fixed' if "تومان" in update.message.text else 'percent'
    await safe_reply(update, "مقدار:", CANCEL_KB)
    return WAIT_BULK_VALUE

async def bulk_value(update, context):
    try: context.user_data['bulk_val'] = float(Utils.to_english(update.message.text))
    except: return WAIT_BULK_VALUE
    await safe_reply(update, "تایید؟", ReplyKeyboardMarkup([["تایید اجرا"], ["انصراف"]], resize_keyboard=True))
    return WAIT_BULK_CONFIRM

async def bulk_confirm(update, context):
    if update.message.text != "تایید اجرا": return await cancel(update, context)
    if await check_is_busy(update, context): return MENU_STATE
    set_busy(context, True)
    msg = await safe_reply(update, "🚀 شروع...")
    cat, val, method = context.user_data['bulk_cat'], context.user_data['bulk_val'], context.user_data['bulk_type']
    count = 0
    try:
        async for products in WCService.fetch_products_generator(cat):
            batch = {"update": []}
            for p in products:
                try:
                    old = float(p.get('regular_price') or 0)
                    if old == 0: continue
                    new = old + val if method == 'fixed' else old + (old * val / 100)
                    batch["update"].append({"id": p['id'], "regular_price": str(int(new))})
                except: continue
            if batch["update"]:
                if await WCService.batch_update(batch):
                    count += len(batch["update"])
                    await safe_edit(context, update.effective_chat.id, msg.message_id, f"♻️ آپدیت: {count}...")
        await safe_edit(context, update.effective_chat.id, msg.message_id, f"✅ تمام. {count} تغییر.")
        Security.log_audit(update.effective_user.id, "BULK", "BATCH", f"{val} {method}")
    finally: set_busy(context, False)
    return MENU_STATE

# --- Remove Bulk ---
async def remove_start(update, context):
    await safe_reply(update, "محدوده:", ReplyKeyboardMarkup([["همه محصولات"], ["انتخاب دسته‌بندی"], ["انصراف"]], resize_keyboard=True))
    return WAIT_REM_SCOPE

async def remove_scope(update, context):
    txt = update.message.text
    if txt == "انصراف": return await cancel(update, context)
    if txt == "همه محصولات": context.user_data['rem_cat'] = None; return await remove_confirm_ask(update, context)
    msg = await safe_reply(update, "⏳...")
    cats = await WCService.get_categories(0)
    await context.bot.delete_message(update.effective_chat.id, msg.message_id)
    context.user_data['cat_map'] = {c['name']: c['id'] for c in cats}
    await safe_reply(update, "دسته:", ReplyKeyboardMarkup([[c['name']] for c in cats], resize_keyboard=True))
    return WAIT_REM_CAT

async def remove_cat(update, context):
    context.user_data['rem_cat'] = context.user_data['cat_map'].get(update.message.text)
    context.user_data['rem_name'] = update.message.text
    return await remove_confirm_ask(update, context)

async def remove_confirm_ask(update, context):
    await safe_reply(update, f"⚠️ تایید حذف؟", ReplyKeyboardMarkup([["تایید حذف"], ["انصراف"]], resize_keyboard=True))
    return WAIT_REM_CONFIRM

async def remove_confirm(update, context):
    if update.message.text != "تایید حذف": return await cancel(update, context)
    if await check_is_busy(update, context): return MENU_STATE
    set_busy(context, True)
    msg = await safe_reply(update, "🚀 پاکسازی...")
    count = 0
    try:
        async for products in WCService.fetch_products_generator(context.user_data['rem_cat']):
            batch = {"update": [{"id": p['id'], "regular_price": "", "sale_price": ""} for p in products]}
            if await WCService.batch_update(batch):
                count += len(batch["update"])
                await safe_edit(context, update.effective_chat.id, msg.message_id, f"🗑 حذف: {count}...")
        await safe_edit(context, update.effective_chat.id, msg.message_id, f"✅ تمام. {count} حذف شد.")
        Security.log_audit(update.effective_user.id, "BULK_REMOVE", "BATCH", context.user_data['rem_name'])
    finally: set_busy(context, False)
    return MENU_STATE

# --- Export Module (Excel Only + Graphics + Stock Filter) ---

async def export_start(update, context):
    # پرسش در مورد فیلتر موجودی به جای فرمت فایل
    btns = [["همه محصولات"], ["فقط موجودها", "فقط ناموجودها"], ["انصراف"]]
    await safe_reply(update, "فیلتر بر اساس موجودی؟", ReplyKeyboardMarkup(btns, resize_keyboard=True))
    return WAIT_EXP_TYPE

async def export_type(update, context):
    txt = update.message.text
    if txt == "انصراف": return await cancel(update, context)
    
    # تعیین فیلتر
    stock_filter = None
    if "فقط موجودها" in txt: stock_filter = "instock"
    elif "فقط ناموجودها" in txt: stock_filter = "outofstock"
    
    context.user_data['exp_filter'] = stock_filter
    
    await safe_reply(update, "محدوده خروجی؟", ReplyKeyboardMarkup([["کل محصولات"], ["انتخاب دسته"], ["انصراف"]], resize_keyboard=True))
    return WAIT_EXP_SCOPE

async def export_scope_handler(update, context):
    txt = update.message.text
    if txt == "انصراف": return await cancel(update, context)
    
    if txt == "کل محصولات":
        return await perform_export(update, context, None)
    
    elif txt == "انتخاب دسته":
        msg = await safe_reply(update, "⏳ دریافت دسته‌ها...")
        cats = await WCService.get_categories(0)
        await context.bot.delete_message(update.effective_chat.id, msg.message_id)
        
        context.user_data['cat_map'] = {c['name']: c['id'] for c in cats}
        btns = [[c['name']] for c in cats] + [["انصراف"]]
        await safe_reply(update, "📂 کدام دسته؟", ReplyKeyboardMarkup(btns, resize_keyboard=True))
        return WAIT_EXP_CAT
    
    return WAIT_EXP_SCOPE

async def export_cat_handler(update, context):
    txt = update.message.text
    if txt == "انصراف": return await cancel(update, context)
    
    cat_id = context.user_data['cat_map'].get(txt)
    if not cat_id:
        await safe_reply(update, "⚠️ نامعتبر.")
        return WAIT_EXP_CAT
        
    return await perform_export(update, context, cat_id)

async def perform_export(update, context, cat_id):
    if await check_is_busy(update, context): return MENU_STATE
    set_busy(context, True)
    
    stock_filter = context.user_data.get('exp_filter')
    # این خط را ساده کردم تا متن فیلتر به کاربر نمایش داده نشود
    msg = await safe_reply(update, "⏳ دریافت و طراحی گزارش گرافیکی...")
    
    data = []
    
    try:
        async for products in WCService.fetch_products_generator(cat_id, stock_status=stock_filter):
            for p in products:
                cats = p.get('categories', [])
                cat_name = cats[0]['name'] if cats else "بدون دسته"
                
                data.append({
                    "دسته بندی": cat_name,
                    "نام محصول": p['name'], 
                    "SKU": p['sku'], 
                    "قیمت (تومان)": p.get('regular_price') or 0, 
                    "موجودی": p.get('stock_quantity') or 0
                })
            
            if len(data) % 50 == 0:
                await safe_edit(context, update.effective_chat.id, msg.message_id, f"📥 جمع‌آوری: {len(data)}...")
    
        if not data:
            await safe_reply(update, "📭 با این فیلتر محصولی یافت نشد.")
            set_busy(context, False)
            return MENU_STATE

        await safe_edit(context, update.effective_chat.id, msg.message_id, "🎨 در حال طراحی گرافیک اکسل...")
        data.sort(key=lambda x: x['دسته بندی'])

        # --- ساخت فایل گرافیکی ---
        def _make_styled_file():
            try:
                # تغییر نام فایل به یک نام ثابت برای جلوگیری از نمایش کلمه instock در اسم فایل
                filename = "Sobhan_List.xlsx"
                
                # 1. نوشتن داده خام (شروع از ردیف 5)
                df = pd.DataFrame(data)
                df.to_excel(filename, index=False, startrow=4)
                
                # 2. باز کردن برای استایل دهی
                wb = load_workbook(filename)
                ws = wb.active
                ws.sheet_view.rightToLeft = True 
                
                # استایل‌ها
                header_font = Font(name='B Titr', size=18, bold=True, color="FFFFFF")
                subheader_font = Font(name='B Nazanin', size=12, bold=True)
                warning_font = Font(name='B Nazanin', size=12, bold=True, color="FF0000") # فونت قرمز برای هشدار
                th_font = Font(name='Tahoma', size=10, bold=True, color="FFFFFF")
                
                # --- هدر 1: نام فروشگاه (A1:E2) ---
                ws.merge_cells('A1:E2') 
                title_cell = ws['A1']
                title_cell.value = "فروشگاه سبحان رایانه"
                title_cell.font = header_font
                title_cell.alignment = Alignment(horizontal='center', vertical='center')
                title_cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # آبی تیره
                
                # --- هدر 2: تاریخ شمسی (A3:E3) ---
                shamsi_date = jdatetime.date.today().strftime("%Y/%m/%d")
                ws.merge_cells('A3:E3')
                date_cell = ws['A3']
                date_cell.value = f"تاریخ گزارش: {shamsi_date}"
                date_cell.font = subheader_font
                date_cell.alignment = Alignment(horizontal='center', vertical='center')
                date_cell.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid") # آبی روشن

                # --- هدر 3: هشدار نوسان قیمت (A4:E4) ---
                ws.merge_cells('A4:E4')
                warn_cell = ws['A4']
                warn_cell.value = "با توجه به نوسانات ارز برای استعلام دقیق قیمت با شماره های زیر تماس بگیرید :\n۰۹۱۳۸۰۰۱۵۹۷\n۰۳۱۳۶۶۳۱۵۹۷"
                warn_cell.font = warning_font
                warn_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True) # wrap_text برای چندخطی شدن
                ws.row_dimensions[4].height = 65 # افزایش ارتفاع ردیف هشدار
                
                # --- هدر 4: سرستون‌های جدول (ردیف 5) ---
                for col in range(1, 6):
                    cell = ws.cell(row=5, column=col)
                    cell.font = th_font
                    cell.fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid") # سرمه‌ای
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    
                # عرض ستون‌ها
                ws.column_dimensions['A'].width = 25 
                ws.column_dimensions['B'].width = 50 
                ws.column_dimensions['C'].width = 15 
                ws.column_dimensions['D'].width = 20 
                ws.column_dimensions['E'].width = 10 

                wb.save(filename)
                return filename
            except Exception as e:
                config.logger.error(f"Excel Style Error: {e}")
                return None

        fname = await asyncio.to_thread(_make_styled_file)
        
        if fname:
            # کپشن بدون وضعیت
            caption = f"✅ لیست موجودی سبحان رایانه\n📅 تاریخ: {jdatetime.date.today().strftime('%Y/%m/%d')}"
            await update.message.reply_document(open(fname, 'rb'), caption=caption)
            os.remove(fname)
        else:
            await safe_reply(update, "❌ خطا در ساخت فایل.")
            
        try: await context.bot.delete_message(update.effective_chat.id, msg.message_id)
        except: pass
        
    finally:
        set_busy(context, False)
        
    return MENU_STATE